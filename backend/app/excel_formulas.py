"""
excel_formulas.py
-----------------
For every answer or change the agent produces, describes how to get the
same result in Microsoft Excel by hand - a formula to paste, the menu
steps, or both.

The point is that a person should never be locked into this tool to act
on what it found. If the assistant says 138 rows are missing a Category,
the drawer shows a formula that returns 138 in their own copy of the
sheet.

How the formulas stay honest
----------------------------
They are built from the same rules the Python code applies, not written
as plausible-looking approximations:

  - Excel's "=" ignores case; the graph matches names exactly. So name
    comparisons use EXACT(), or "Sakal" and "SAKAL" would be merged.
  - The graph skips a row unless both the advertiser and the publication
    are present, so graph formulas carry the same (B<>"") condition.
  - "Missing" here means empty OR only spaces, so gap counts use
    TRIM(...)="" rather than COUNTBLANK, which ignores cells holding a space.
  - For edits, the cell list comes from diffing the data before and after
    the change, so it names exactly the cells that moved.

Each formula is checked by tests/verify_excel_formulas.py, which writes it
into a real workbook, has Excel calculate it, and compares the result
with what the Python code returned.

The one thing that cannot be derived this way is a free-form table query,
because the model writes that expression. Its Excel equivalent comes from
the model too, and is labelled as unchecked in the response.

Sheet assumptions, stated in every result: headers in row 1, data from
row 2, columns in the original order - which is exactly the layout of
the uploaded file and of the file this app exports.
"""

from __future__ import annotations

import pandas as pd
from openpyxl.utils import get_column_letter

LAYOUT_NOTE = "Assumes the sheet layout of your file: headers in row 1, data starting in row 2."
DYNAMIC_ARRAY_NOTE = "Uses FILTER and UNIQUE, which need Excel 365 or Excel 2021 or later."
SEPARATOR_NOTE = "If your Excel uses semicolons between arguments, replace each comma with a semicolon."

# The Name Box rejects entries much past 255 characters.
_NAME_BOX_LIMIT = 250
# Past this many individual cells, typing addresses stops being a sensible instruction.
_MAX_LISTED_CELLS = 400


# ------------------------------------------------------------------ helpers

def _q(text: str) -> str:
    """A string as an Excel literal, with embedded quotes doubled."""
    return '"' + str(text).replace('"', '""') + '"'


def col_letter(df: pd.DataFrame, column: str) -> str | None:
    if column not in df.columns:
        return None
    return get_column_letter(df.columns.get_loc(column) + 1)


def excel_row(index: int) -> int:
    """Dataframe index -> sheet row. Row 1 holds the headers."""
    return int(index) + 2


def _last_row(df: pd.DataFrame) -> int:
    return excel_row(df.index.max()) if len(df) else 2


def _rng(df: pd.DataFrame, column: str) -> str:
    letter = col_letter(df, column)
    return f"${letter}$2:${letter}${_last_row(df)}"


def _formula(label: str, formula: str, result=None) -> dict:
    out = {"label": label, "formula": formula}
    if result is not None:
        out["expected"] = result
    return out


def _name_box_chunks(cells: list[str]) -> list[str]:
    chunks, current = [], ""
    for cell in cells:
        candidate = f"{current},{cell}" if current else cell
        if len(candidate) > _NAME_BOX_LIMIT:
            chunks.append(current)
            current = cell
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


def sheet_layout(df: pd.DataFrame) -> str:
    """Column letters for the model, so any formula it suggests points at real columns."""
    cols = ", ".join(f"{get_column_letter(i + 1)}={c}" for i, c in enumerate(df.columns))
    return f"Sheet layout: headers in row 1, data in rows 2 to {_last_row(df)}. Columns: {cols}."


# ------------------------------------------------------------------ reads

def _missing_value(df: pd.DataFrame, column: str, count: int) -> dict:
    r = _rng(df, column)
    letter = col_letter(df, column)
    return {
        "title": f"Count the blanks in {column}",
        "formulas": [_formula(f"Rows with no {column}", f'=SUMPRODUCT(--(TRIM({r})=""))', count)],
        "steps": [
            "Click any cell in the header row, then Data › Filter.",
            f"Open the dropdown on {column} (column {letter}).",
            "Untick Select All, tick (Blanks), and press OK.",
            "The status bar shows how many rows are left.",
        ],
        "notes": [
            "Counts cells that are empty or hold only spaces, which is what the check treats as missing.",
        ],
    }


def _name_mismatch(df: pd.DataFrame) -> dict | None:
    raw, ai = "AdvertiserName", "Advertiser Name by AI"
    if raw not in df.columns or ai not in df.columns:
        return None
    both = df.dropna(subset=[raw, ai])
    rows = int((both[raw].str.strip().str.lower() != both[ai].str.strip().str.lower()).sum())
    a, b = _rng(df, raw), _rng(df, ai)
    return {
        "title": "Rows where the two advertiser-name columns disagree",
        "formulas": [_formula(
            "Rows with a mismatch",
            f'=SUMPRODUCT(({a}<>"")*({b}<>"")*(LOWER(TRIM({a}))<>LOWER(TRIM({b}))))',
            rows,
        )],
        "steps": [
            f"In an empty column, enter =LOWER(TRIM({col_letter(df, raw)}2))<>LOWER(TRIM({col_letter(df, ai)}2)) in row 2 and fill it down.",
            "Filter that column to TRUE to list the disagreeing rows.",
        ],
        "notes": ["Counts rows. The Problems panel groups them by raw name, so its count is lower."],
    }


def for_detect_issues(df: pd.DataFrame, arguments: dict, result: dict) -> dict | None:
    items = []
    seen_mismatch = False
    for issue in result.get("issues", []):
        kind, scope = issue.get("type"), issue.get("scope", "")
        if kind == "missing_value" and scope.startswith("column:"):
            column = scope.split(":", 1)[1]
            if column in df.columns:
                missing = int((df[column].isna() | (df[column].astype(str).str.strip() == "")).sum())
                items.append(_missing_value(df, column, missing))
        elif kind == "name_mismatch" and not seen_mismatch:
            seen_mismatch = True
            m = _name_mismatch(df)
            if m:
                items.append(m)

    if any(i.get("type") == "near_duplicate_advertiser" for i in result.get("issues", [])):
        items.append({
            "title": "Similar advertiser names",
            "formulas": [],
            "steps": [
                "Excel has no formula for similarity between two names.",
                "Data › Get Data › From Table/Range opens Power Query, where Merge Queries with "
                "'Use fuzzy matching' finds the same kind of near-duplicates.",
            ],
            "notes": ["This check compares spellings, so it has no exact formula equivalent."],
        })
    return _wrap(items) if items else None


def for_entity_summary(df: pd.DataFrame, arguments: dict, result: dict) -> dict | None:
    """Counts and distinct-value lists for one subject, on the user's sheet."""
    if "error" in result:
        return None
    name = arguments.get("entity", "")
    subject_col = result.get("subject_column")
    if not subject_col or subject_col not in df.columns:
        return None

    subj = _rng(df, subject_col)
    match = f"EXACT({subj},{_q(name)})"
    formulas = [_formula(f"Rows for this {subject_col}", f"=SUMPRODUCT(--{match})", result.get("row_count"))]

    for column, values in (result.get("linked") or {}).items():
        if column not in df.columns:
            continue
        r = _rng(df, column)
        formulas.append(_formula(
            f"Distinct {column}",
            f'=UNIQUE(FILTER({r},{match}*({r}<>""),"none"))',
            f"{len(values)} value(s)",
        ))

    for flag, count in (result.get("flags") or {}).items():
        if flag in df.columns:
            formulas.append(_formula(f"Rows flagged in {flag}",
                                     f'=SUMPRODUCT({match}*(TRIM({_rng(df, flag)})<>""))', count))

    return _wrap([{
        "title": f"Summary of {name}",
        "formulas": formulas,
        "steps": [
            "Paste each formula into an empty cell.",
            "The list formulas spill downward, so leave the cells below them empty.",
        ],
        "notes": [
            "Names are matched exactly, including capitals, as the assistant does.",
            DYNAMIC_ARRAY_NOTE,
        ],
    }])


def for_find_by_attributes(df: pd.DataFrame, arguments: dict, result: dict) -> dict | None:
    """Subjects matching every criterion, where the criteria may be
    satisfied by different rows of the same subject."""
    criteria = {k: v for k, v in (arguments.get("criteria") or {}).items() if k in df.columns}
    if len(criteria) < 1:
        return None
    subject_col = None
    for name, role in (df.attrs.get("roles") or {}).items():
        if role == "subject":
            subject_col = name
    subject_col = subject_col or _first_subject(df)
    if subject_col is None:
        return None

    subj = _rng(df, subject_col)
    lists, names = [], []
    for i, (column, value) in enumerate(criteria.items()):
        var = f"set{i + 1}"
        names.append(var)
        lists.append(f'{var},UNIQUE(FILTER(names,valid*EXACT({_rng(df, column)},{_q(value)}),""))')

    keep = names[0]
    for other in names[1:]:
        keep = f"FILTER({keep},BYROW({keep},LAMBDA(n,OR(EXACT(n,{other})))),\"none\")"

    formula = (f'=LET(names,{subj},valid,({subj}<>""),' + ",".join(lists) + f",{keep})")
    return _wrap([{
        "title": "Subjects matching " + " and ".join(f"{c}={v}" for c, v in criteria.items()),
        "formulas": [_formula("Matching values", formula, f"{result.get('count', 0)} match(es)")],
        "steps": [
            "Paste into an empty cell; the list spills downward.",
            "Wrap it in ROWS( ... ) to get just the count.",
        ],
        "notes": [
            "Each criterion is matched separately and the lists intersected, because a subject can "
            "satisfy them on different rows. A single FILTER with both conditions would answer a "
            "different question.",
            "Needs Excel 365, for LET, BYROW and LAMBDA.",
        ],
    }])


def _first_subject(df: pd.DataFrame) -> str | None:
    from app.schema_profile import infer_mapping
    try:
        return infer_mapping(df).get("subject")
    except Exception:  # noqa: BLE001
        return None


def for_tabular_query(df: pd.DataFrame, arguments: dict, result: dict) -> dict | None:
    suggested = (arguments.get("excel_formula") or "").strip()
    if not suggested or "error" in result:
        return None
    if not suggested.startswith("="):
        suggested = "=" + suggested
    return _wrap([{
        "title": "Equivalent formula",
        "formulas": [_formula("Suggested formula", suggested)],
        "steps": ["Paste into an empty cell and compare the result with the answer."],
        "notes": [
            "Written by the assistant alongside its query and not checked automatically. "
            "Compare its result before relying on it.",
        ],
        "unverified": True,
    }])


READ_BUILDERS = {
    "detect_issues": for_detect_issues,
    "get_entity_summary": for_entity_summary,
    "find_by_attributes": for_find_by_attributes,
    "tabular_query": for_tabular_query,
}


def for_read(name: str, df: pd.DataFrame, arguments: dict, result: dict) -> dict | None:
    builder = READ_BUILDERS.get(name)
    if builder is None or not isinstance(result, dict):
        return None
    try:
        return builder(df, arguments, result)
    except Exception:  # noqa: BLE001 - a missing formula must never break an answer
        return None


# ------------------------------------------------------------------ edits

def _changed_rows(before: pd.DataFrame, after: pd.DataFrame, column: str) -> list[int]:
    if column not in after.columns:
        return []
    old = before[column] if column in before.columns else pd.Series("", index=after.index)
    blank_old = old.isna() | (old.astype(str) == "")
    blank_new = after[column].isna() | (after[column].astype(str) == "")
    differs = (old.astype(str) != after[column].astype(str)) & ~(blank_old & blank_new)
    return [int(i) for i in after.index[differs]]


def _cell_steps(letter: str, rows: list[int], value: str, what: str) -> tuple[list[str], list[str]]:
    cells = [f"{letter}{excel_row(r)}" for r in rows]
    chunks = _name_box_chunks(cells)
    steps = [
        "Click the Name Box, left of the formula bar.",
        f"Type the cell list below and press Enter. This selects all {len(cells)} {what} at once.",
        f"Type {value} and press Ctrl+Enter, which fills every selected cell.",
    ]
    if len(chunks) > 1:
        steps.append(f"The list is split into {len(chunks)} parts because the Name Box has a length limit. Repeat for each part.")
    return steps, chunks


def for_edit(operation: str, arguments: dict, before: pd.DataFrame, after: pd.DataFrame) -> dict | None:
    try:
        return _for_edit(operation, arguments, before, after)
    except Exception:  # noqa: BLE001
        return None


def _for_edit(operation: str, arguments: dict, before: pd.DataFrame, after: pd.DataFrame) -> dict | None:
    value = arguments.get("value") or ("flagged" if operation == "flag_rows" else "")

    if operation == "fill_missing_value_bulk":
        column = arguments["column"]
        letter = col_letter(after, column)
        rows = _changed_rows(before, after, column)
        r = _rng(before, column)
        first = excel_row(before.index.min())
        return _wrap([{
            "title": f"Fill every blank in {column} with {value}",
            "formulas": [
                _formula("Blanks before the change", f'=SUMPRODUCT(--(TRIM({r})=""))', len(rows)),
                _formula(
                    f"Helper column (enter in row {first}, fill down)",
                    f'=IF(TRIM({letter}{first})="",{_q(value)},{letter}{first})',
                ),
            ],
            "steps": [
                f"Select the data in column {letter}, {letter}{first}:{letter}{_last_row(before)}.",
                "Home › Find & Select › Go To Special › Blanks › OK.",
                f"Type {value} and press Ctrl+Enter.",
                "Run the blank count again. It should now return 0.",
            ],
            "notes": [
                "Go To Special skips cells that contain only a space. If the count is not 0 afterwards, "
                "use the helper column instead, then copy it and Paste Values over the original.",
            ],
        }], changed=len(rows))

    if operation in ("fill_missing_value", "standardize_value"):
        column = arguments["column"]
        letter = col_letter(after, column)
        rows = _changed_rows(before, after, column)
        if not rows:
            return None
        verb = "Fill" if operation == "fill_missing_value" else "Set"
        title = f"{verb} {column} to {value} in {len(rows)} row(s)"
        if len(rows) > _MAX_LISTED_CELLS:
            return _wrap([{
                "title": title,
                "formulas": [],
                "steps": [
                    f"This change touched {len(rows)} cells, too many to select by address.",
                    "Export this version from the app and open it in Excel instead.",
                ],
                "notes": [],
            }], changed=len(rows))
        steps, chunks = _cell_steps(letter, rows, value, "cells")
        return _wrap([{
            "title": title,
            "formulas": [],
            "cells": chunks,
            "steps": steps,
            "notes": ["The cells listed are exactly the ones this change altered."],
        }], changed=len(rows))

    if operation == "flag_rows":
        flag_column = arguments.get("flag_column", "review_flag")
        letter = col_letter(after, flag_column)
        rows = _changed_rows(before, after, flag_column)
        if not rows:
            return None
        new_column = flag_column not in before.columns
        steps = []
        if new_column:
            steps.append(f"Type {flag_column} into {letter}1 to add the column.")
        if len(rows) > _MAX_LISTED_CELLS:
            steps.append("Too many rows to select by address; export this version from the app instead.")
            return _wrap([{"title": f"Flag {len(rows)} row(s)", "formulas": [], "steps": steps, "notes": []}],
                         changed=len(rows))
        cell_steps, chunks = _cell_steps(letter, rows, value, "rows")
        return _wrap([{
            "title": f"Flag {len(rows)} row(s) in {flag_column}",
            "formulas": [],
            "cells": chunks,
            "steps": steps + cell_steps,
            "notes": ["Only the flag column changes; no original values are touched."],
        }], changed=len(rows))

    return None


def _wrap(items: list[dict], changed: int | None = None) -> dict:
    out = {"items": items, "notes": [LAYOUT_NOTE, SEPARATOR_NOTE]}
    if changed is not None:
        out["changed_cells"] = changed
    return out
