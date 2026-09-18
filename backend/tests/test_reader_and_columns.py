"""Locating the table in a sheet, and adding columns."""

import io

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.config import settings
from app.excel_reader import read_table
from app.main import app
from app.safe_editor import EditValidationError, add_column

client = TestClient(app)


def _xlsx(rows, start_row=1, start_col=1, sheets=None) -> bytes:
    """Writes `rows` with the top-left cell at (start_row, start_col)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for name, content in (sheets or {}).items():
        extra = wb.create_sheet(name, 0)
        for r, row in enumerate(content, start=1):
            for c, v in enumerate(row, start=1):
                extra.cell(r, c, v)
    for r, row in enumerate(rows, start=start_row):
        for c, v in enumerate(row, start=start_col):
            if v is not None:
                ws.cell(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


TABLE = [
    ["Name", "City", "Amount"],
    ["Asha", "Pune", 10],
    ["Ravi", None, 20],
    ["Meera", "Mumbai", 30],
]


def test_plain_table_at_a1_is_unchanged():
    df, where = read_table(_xlsx(TABLE))
    assert list(df.columns) == ["Name", "City", "Amount"]
    assert len(df) == 3
    assert where.origin == "A1" and not where.moved
    assert df["Amount"].tolist() == [10, 20, 30]


def test_table_starting_at_c3_with_title_above():
    rows = [["Quarterly report", None, None], [None, None, None]] + TABLE
    df, where = read_table(_xlsx(rows, start_row=1, start_col=3))
    assert list(df.columns) == ["Name", "City", "Amount"]
    assert where.origin == "C3" and where.moved
    assert df["Name"].tolist() == ["Asha", "Ravi", "Meera"]
    assert pd.isna(df.at[1, "City"])          # blank stays blank, not 0
    assert df["Amount"].tolist() == [10, 20, 30]


def test_headings_on_row_four_under_notes():
    rows = [["Sales export"], ["Generated 2026-09-18"], [None]] + TABLE
    df, where = read_table(_xlsx(rows))
    assert list(df.columns) == ["Name", "City", "Amount"]
    assert where.header_row == 4


def test_blank_heading_named_after_its_letter():
    rows = [["Name", None, "Amount"], ["Asha", "x", 10], ["Ravi", "y", 20]]
    df, _ = read_table(_xlsx(rows))
    assert list(df.columns) == ["Name", "Column B", "Amount"]


def test_numeric_headings_become_text():
    rows = [["Region", 2024, 2025], ["North", 1, 2], ["South", 3, 4]]
    df, _ = read_table(_xlsx(rows))
    assert list(df.columns) == ["Region", "2024", "2025"]


def test_empty_cover_sheet_is_skipped():
    df, where = read_table(_xlsx(TABLE, sheets={"Cover": []}))
    assert where.sheet == "Data" and where.moved
    assert len(df) == 3


def test_empty_spacer_rows_dropped_values_kept():
    rows = TABLE[:2] + [[None, None, None]] + TABLE[2:]
    df, _ = read_table(_xlsx(rows))
    assert len(df) == 3
    assert list(df.index) == [0, 1, 2]


def test_sample_file_reads_as_before():
    contents = settings.resolved_sample_file.read_bytes()
    df, where = read_table(contents)
    assert df.shape == (11275, 16)
    pd.testing.assert_frame_equal(df, pd.read_excel(io.BytesIO(contents)))
    assert not where.moved


# ------------------------------------------------------------------ add column

DF = pd.DataFrame({"Name": ["Asha", "Ravi"], "City": ["Pune", None]})


def test_add_empty_column_at_end():
    new, _ = add_column(DF, "Notes")
    assert list(new.columns) == ["Name", "City", "Notes"]
    assert new["Notes"].isna().all()
    assert list(DF.columns) == ["Name", "City"]          # original untouched


def test_add_column_with_value_after_given_column():
    new, preview = add_column(DF, "Status", value="open", after_column="Name")
    assert list(new.columns) == ["Name", "Status", "City"]
    assert new["Status"].tolist() == ["open", "open"]
    assert preview[0]["after"] == "open"


def test_add_column_copied_from_another():
    new, _ = add_column(DF, "City copy", source_column="City")
    assert new["City copy"].tolist()[0] == "Pune" and pd.isna(new["City copy"].tolist()[1])


@pytest.mark.parametrize("kwargs, message", [
    ({"new_column": "  "}, "needs a name"),
    ({"new_column": "city"}, "already exists"),
    ({"new_column": "X", "value": "a", "source_column": "City"}, "not both"),
    ({"new_column": "X", "source_column": "Nope"}, "does not exist"),
    ({"new_column": "X", "after_column": "Nope"}, "does not exist"),
])
def test_add_column_refusals(kwargs, message):
    with pytest.raises(EditValidationError, match=message):
        add_column(DF, **kwargs)


def test_columns_endpoint_creates_audited_version():
    session_id = client.post("/sample").json()["session_id"]
    r = client.post("/columns", json={"session_id": session_id, "name": "Reviewed", "value": "no",
                                      "after_column": "Category"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["current_version"] == 1 and body["rows_affected"] == 11275
    assert body["excel"]["items"][0]["title"] == "Add column Reviewed"

    page = client.get(f"/data/{session_id}", params={"limit": 1}).json()
    cols = page["columns"]
    assert cols[cols.index("Category") + 1] == "Reviewed"

    diff = client.get(f"/diff/{session_id}/1").json()
    assert diff["added_columns"] == ["Reviewed"]

    hist = client.get(f"/history/{session_id}").json()
    assert hist["audit_log"][-1]["operation"] == "add_column"

    dup = client.post("/columns", json={"session_id": session_id, "name": "Reviewed"})
    assert dup.status_code == 400


def test_upload_reports_where_table_was_found():
    rows = [["Report"], [None]] + TABLE
    r = client.post("/upload", files={"file": ("r.xlsx", _xlsx(rows, start_col=2),
                                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["columns"] == ["Name", "City", "Amount"]
    assert "B3" in body["notice"]
